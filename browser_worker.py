# -*- coding: utf-8 -*-
"""
浏览器工作线程：独立线程持有 Playwright 持久化上下文（系统 Edge），
通过命令队列与 GUI 主线程通信，避免阻塞 tkinter 主循环。
"""
import os, time, threading, queue
from datetime import datetime
from playwright.sync_api import sync_playwright
from openpyxl.utils import get_column_letter

from adapters import get_adapter, rewrite_img_url, FIELD_LABELS


class BrowserWorker(threading.Thread):
    def __init__(self, profile_dir, log_cb=None):
        super().__init__(daemon=True)
        self.profile_dir = profile_dir
        self.cmd_q = queue.Queue()
        self.result_q = queue.Queue()
        self.log_cb = log_cb
        self._ctx = None
        self._page = None
        self._stop = False

    def log(self, msg):
        if self.log_cb:
            self.log_cb(msg)

    # ---------- GUI 侧调用（投递命令） ----------
    def post(self, cmd_type, **kw):
        self.cmd_q.put({"type": cmd_type, **kw})

    def open_url(self, url):
        self.post("open_url", url=url)

    def identify(self):
        self.post("identify")

    def scrape(self, max_records=0, fields=None, out_dir=None):
        self.post("scrape", max_records=max_records, fields=fields or [], out_dir=out_dir)

    def close_browser(self):
        self._stop = True
        self.post("close")

    # ---------- 线程主体 ----------
    def run(self):
        with sync_playwright() as p:
            self.log("正在启动内置浏览器…")
            os.makedirs(self.profile_dir, exist_ok=True)
            self._ctx = p.chromium.launch_persistent_context(
                user_data_dir=self.profile_dir,
                channel="msedge",
                headless=False,
                viewport={"width": 1366, "height": 850},
                args=["--no-first-run", "--no-default-browser-check",
                      "--start-maximized"],
            )
            self._page = self._ctx.pages[0] if self._ctx.pages else self._ctx.new_page()
            self.log("浏览器已就绪（登录态持久保存在程序目录 profiles\\default）")
            self.result_q.put({"type": "ready"})
            while not self._stop:
                try:
                    cmd = self.cmd_q.get(timeout=0.5)
                except queue.Empty:
                    continue
                try:
                    self._dispatch(cmd)
                except Exception as e:
                    self.log(f"错误: {e}")
                    self.result_q.put({"type": "error", "msg": str(e)})
            try:
                self._ctx.close()
            except Exception:
                pass

    def _dispatch(self, cmd):
        t = cmd["type"]
        if t == "open_url":
            url = cmd.get("url", "").strip()
            if not url.startswith("http"):
                url = "https://" + url
            self.log(f"打开: {url}")
            self._page.goto(url, timeout=45000, wait_until="domcontentloaded")
            self.log("页面已打开")
            self.result_q.put({"type": "page_loaded", "url": self._page.url,
                               "title": self._page.title()})
        elif t == "identify":
            self._identify()
        elif t == "scrape":
            self._scrape(cmd)
        elif t == "close":
            self._stop = True

    def _identify(self):
        url = self._page.url
        adapter = get_adapter(url)
        self.log(f"识别平台: {adapter.name}")
        # 滚动探测数量：滚几轮看目标元素数
        count = self._probe_count(adapter)
        self.result_q.put({"type": "identified", "url": url,
                           "adapter": adapter.name,
                           "fields": adapter.fields, "count": count})

    def _probe_count(self, adapter, rounds=6):
        last = -1
        for _ in range(rounds):
            try:
                n = self._page.evaluate(adapter.count_js())
            except Exception:
                return 0
            self._page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
            time.sleep(0.6)
            if n == last:
                break
            last = n
        self._page.evaluate("window.scrollTo(0, 0)")
        return last if last > 0 else 0

    def _scroll_load(self, adapter, max_records):
        """滚动加载：拼多多按店铺卡片计数，通用适配器按高度"""
        stable = 0
        for i in range(150):
            try:
                n = self._page.evaluate(adapter.count_js())
            except Exception:
                break
            self._page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
            time.sleep(0.7)
            if max_records and isinstance(n, int) and n >= max_records:
                self.log(f"已加载到目标数量 {n}，停止滚动")
                break
            if n == stable:
                break
            stable = n
            if i % 5 == 0:
                self.log(f"滚动加载中… 当前 {n}")

    def _scrape(self, cmd):
        max_records = int(cmd.get("max_records") or 0)
        want_fields = cmd.get("fields") or []
        out_dir = cmd.get("out_dir") or "."
        url = self._page.url
        adapter = get_adapter(url)
        self.log(f"开始抓取 [{adapter.name}] 页面: {url[:80]}")

        self._scroll_load(adapter, max_records)
        time.sleep(1)

        rows = self._page.evaluate(adapter.extract_js)
        self.log(f"提取到 {len(rows)} 条记录")
        if not rows:
            self.result_q.put({"type": "scrape_done", "rows": [], "path": None,
                               "error": "未提取到数据，请确认已登录且页面已加载完成"})
            return

        # 过滤字段 + 数量截断
        limit = max_records if max_records > 0 else len(rows)
        rows = rows[:limit]
        # 图片下载
        req_ctx = self._page.context.request
        for i, r in enumerate(rows):
            if r.get("img") and "img" in want_fields:
                try:
                    u = rewrite_img_url(r["img"], url)
                    resp = req_ctx.get(u, timeout=20000)
                    if not resp.ok:
                        resp = req_ctx.get(r["img"], timeout=20000)
                    if resp.ok:
                        body = resp.body()
                        ext = "jpg"
                        if body[:8] == b"\x89PNG\r\n\x1a\n":
                            ext = "png"
                        elif body[:4] == b"RIFF" and body[8:12] == b"WEBP":
                            ext = "webp"
                        img_dir = os.path.join(out_dir, "imgs")
                        os.makedirs(img_dir, exist_ok=True)
                        fpath = os.path.join(img_dir, f"img_{i+1:04d}.{ext}")
                        with open(fpath, "wb") as f:
                            f.write(body)
                        r["img_local"] = fpath
                    else:
                        r["img_local"] = None
                except Exception as e:
                    r["img_local"] = None
            else:
                r["img_local"] = None
            if i % 10 == 0:
                self.log(f"图片 {i}/{len(rows)}")

        path = self._export_excel(rows, want_fields, out_dir, url)
        self.log(f"导出完成: {path}")
        self.result_q.put({"type": "scrape_done", "rows": rows,
                           "path": path, "error": None})

    def _export_excel(self, rows, want_fields, out_dir, url):
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

        labels = {"序号": "序号"}
        for f in want_fields:
            labels[f] = FIELD_LABELS.get(f, f)

        wb = Workbook()
        ws = wb.active
        ws.title = "抓取结果"
        header_fill = PatternFill("solid", fgColor="2F5B7C")  # 深蓝灰(柔和)
        header_font = Font(bold=True, color="FFFFFF", size=11)
        thin = Side(style="thin", color="C0C0C0")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        col_idx = {}
        for c, f in enumerate(["序号"] + want_fields, 1):
            col_idx[f] = c
            cell = ws.cell(row=1, column=c, value=labels.get(f, f))
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        widths = {"序号": 6, "shop": 22, "name": 45, "spec": 28, "price": 10,
                  "qty": 8, "paid": 13, "status": 14, "img": 14}
        for f in ["序号"] + want_fields:
            ws.column_dimensions[get_column_letter(col_idx[f])].width = widths.get(f, 15)

        # 数量/金额字段写入为数字，便于 Excel 统计求和
        numeric_int = {"qty"}
        numeric_float = {"price", "paid"}

        def _to_number(field, raw):
            if raw is None or raw == "":
                return ""
            if isinstance(raw, (int, float)):
                return raw
            s = str(raw).strip().replace(",", "").replace("，", "")
            s = s.replace("￥", "").replace("¥", "").replace("元", "")
            s = s.replace("x", "").replace("×", "").replace("X", "").strip()
            if not s:
                return ""
            try:
                if field in numeric_int:
                    return int(float(s))
                return float(s) if "." in s else int(s)
            except (ValueError, TypeError):
                return raw

        img_col = col_idx.get("img")
        for i, r in enumerate(rows, 1):
            row = i + 1
            ws.cell(row=row, column=col_idx["序号"], value=i).border = border
            for f in want_fields:
                if f == "img":
                    continue
                val = r.get(f, "")
                if f in numeric_int or f in numeric_float:
                    val = _to_number(f, val)
                cell = ws.cell(row=row, column=col_idx[f], value=val)
                cell.border = border
                align_kw = {"vertical": "center", "wrap_text": (f in ("name", "spec"))}
                if f in numeric_int or f in numeric_float:
                    align_kw["horizontal"] = "right"
                cell.alignment = Alignment(**align_kw)
            if img_col and r.get("img_local") and os.path.exists(r["img_local"]) \
                    and not r["img_local"].endswith(".webp"):
                try:
                    img = XLImage(r["img_local"])
                    img.width = 56
                    img.height = 56
                    ws.add_image(img, f"{get_column_letter(img_col)}{row}")
                except Exception:
                    pass
            ws.row_dimensions[row].height = 42

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = os.path.join(out_dir, f"抓取结果_{ts}.xlsx")
        wb.save(path)
        return path
